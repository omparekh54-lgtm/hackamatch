import csv, io, hashlib, re, zipfile
from collections import Counter
from openpyxl import load_workbook
from sqlalchemy import select, update
from .models import Attendee, Import, Setting, ImportRecord

FIELDS = {
    "skills": "Skills",
    "goals": "What's your main goal for this hackathon?",
    "stretch": "Do you want to work in your strongest area or stretch into a weaker one?",
    "experience": "Have you done a hackathon before?",
    "schedule": "When are you most productive during the event?",
    "conflict": "How do you handle disagreements?",
    "leadership": "Do you naturally take the lead in groups?",
    "hobbies": "Hobbies",
    "description": "Describe yourself in one word",
    "feed": "What's your Instagram feed like",
    "soundtrack": "Background soundtrack while coding",
    "college": "College/University/School Name",
    "year": "Year of study (eg: 1, 2, 3, 4)",
    "availability": "If selected you will be available to come for offline round at mpstme mumbai campus",
}
MAPPINGS = {
    "frontend": ["front end", "frontend", "react", "html", "css"],
    "backend": ["backend", "back end", "django", "fastapi", "node"],
    "fullstack": ["full stack", "full-stack", "fullstack"],
    "ai": ["machine learning", "artificial intelligence", "ml", "ai"],
    "data": ["data science", "analytics"],
    "mobile": ["android", "ios", "flutter", "mobile"],
    "design": ["ui/ux", "ui ux", "design"],
    "hardware": ["iot", "hardware", "arduino"],
    "product": ["business", "product"],
    "presentation": ["presenting", "public speaking"],
    "research": ["research"],
}


def email(value):
    v = str(value or "").strip().lower()
    if len(v) > 254 or not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", v):
        raise ValueError("Invalid email")
    return v


def tokens(s):
    return sorted(
        set(re.findall(r"[a-z0-9]+", str(s or "").lower()))
        - {"and", "the", "a", "in", "of", "to", "i", "my"}
    )


def parts(s):
    return sorted(
        set(x.strip().lower() for x in re.split(r"[,;\n|]+", str(s or "")) if x.strip())
    )


def normalize(raw, mappings=None):
    p = {k: str(raw.get(col) or "").strip() for k, col in FIELDS.items()}
    sk = p["skills"].lower()
    p["skills"] = sorted(
        k
        for k, words in (mappings or MAPPINGS).items()
        if any(re.search(r"(?<!\w)" + re.escape(w) + r"(?!\w)", sk) for w in words)
    )
    if sk and not p["skills"]:
        p["skills"] = ["other"]
    for k in ["goals", "schedule", "soundtrack"]:
        p[k] = parts(p[k])
    for k in ["hobbies", "description", "feed"]:
        p[k] = tokens(p[k])
    p["experience"] = {
        "never": 0,
        "none": 0,
        "first": 0,
        "1–2": 1,
        "1-2": 1,
        "3+": 2,
    }.get(p["experience"].lower())
    p["leadership"] = {"rarely": 0, "sometimes": 1, "usually": 2}.get(
        p["leadership"].lower()
    )
    p["conflict"] = p["conflict"].lower() or None
    p["stretch"] = (
        p["stretch"].lower()
        if p["stretch"].lower() in ["strength", "stretch", "mix"]
        else None
    )
    p["availability"] = {
        "yes": "yes",
        "y": "yes",
        "true": "yes",
        "no": "no",
        "n": "no",
        "false": "no",
    }.get(p["availability"].lower(), "unknown")
    yr = p["year"].lower()
    match = re.search(r"\b([1-4])(?:st|nd|rd|th)?\b", yr)
    p["year"] = (
        int(match[1]) if match else {"sy": 2, "se": 2, "third": 3, "final": 4}.get(yr)
    )
    p["college"] = re.sub(r"\s+", " ", p["college"]).strip().lower() or None
    p["completeness"] = round(
        sum(bool(str(raw.get(FIELDS[k]) or "").strip()) for k in FIELDS)
        / len(FIELDS)
        * 100
    )
    return p


def parse(data, filename):
    if len(data) > 5 * 1024 * 1024:
        raise ValueError("File exceeds 5 MB")
    if filename.lower().endswith(".csv"):
        reader = csv.DictReader(io.StringIO(data.decode("utf-8-sig")))
        headers = reader.fieldnames or []
        rows = list(reader)
    elif filename.lower().endswith(".xlsx"):
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if sum(x.file_size for x in z.infolist()) > 30 * 1024 * 1024:
                raise ValueError("Workbook expands beyond limit")
        ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
        it = ws.iter_rows(values_only=True)
        headers = [str(x or "") for x in next(it)]
        rows = [
            dict(zip(headers, [str(x) if x is not None else "" for x in row]))
            for row in it
            if any(x is not None for x in row)
        ]
    else:
        raise ValueError("Use CSV or XLSX")
    if len(headers) != len(set(headers)):
        raise ValueError("Duplicate headers")
    missing = set(
        ["guest_id", "email", "name", "approval_status", *FIELDS.values()]
    ) - set(headers)
    if missing:
        raise ValueError("Missing columns: " + ", ".join(sorted(missing)))
    if len(rows) > 5000:
        raise ValueError("Maximum 5000 rows per import")
    return headers, rows


def import_data(db, data, filename, actor, preview=False):
    digest = hashlib.sha256(data).hexdigest()
    previous = db.scalar(select(Import).where(Import.digest == digest))
    if previous and not preview:
        return {"id": previous.id, "duplicate_import": True, **previous.summary}
    headers, rows = parse(data, filename)
    setting = db.get(Setting, "skill_mappings")
    mappings = setting.value if setting else MAPPINGS
    valid = []
    rejected = []
    duplicates = []
    ids = set()
    for index, row in enumerate(rows, 2):
        try:
            if None in row or not str(row.get("guest_id") or "").strip():
                raise ValueError("Malformed row or missing guest_id")
            em = email(row.get("email"))
            gid = str(row["guest_id"]).strip()
            if len(gid) > 200:
                raise ValueError("Registration identifier too long")
            if gid in ids:
                duplicates.append(index)
                continue
            ids.add(gid)
            valid.append(
                (
                    gid,
                    em,
                    (
                        str(row["approval_status"]).strip().lower()
                        if str(row["approval_status"]).strip().lower()
                        in ["approved", "declined"]
                        else "requires_review"
                    ),
                    row,
                    normalize(row, mappings),
                )
            )
        except (ValueError, TypeError):
            rejected.append(
                {"row": index, "reason": "Invalid identity or malformed row"}
            )
    # Conflicting rows for one external ID are quarantined, never first-row-wins.
    by_id = {}
    for row in rows:
        gid = str(row.get("guest_id") or "").strip()
        if gid:
            by_id.setdefault(gid, []).append(row)
    conflicting = {
        gid
        for gid, group in by_id.items()
        if len({str(sorted((str(k), str(v)) for k, v in row.items())) for row in group})
        > 1
    }
    valid = [r for r in valid if r[0] not in conflicting]
    rejected.extend(
        {"row": i, "reason": "Conflicting duplicate guest_id requires review"}
        for i, row in enumerate(rows, 2)
        if str(row.get("guest_id") or "").strip() in conflicting
    )
    emails = Counter(r[1] for r in valid)
    summary = {
        "total": len(rows),
        "accepted": len(valid),
        "rejected": rejected,
        "duplicates": duplicates,
        "approved": sum(r[2] == "approved" for r in valid),
        "declined": sum(r[2] == "declined" for r in valid),
        "ambiguous_emails": sum(n > 1 for n in emails.values()),
        "unknown_availability": sum(r[4]["availability"] == "unknown" for r in valid),
        "mapping": {
            "guest_id": "guest_id",
            "email": "email",
            "approval": "approval_status",
            **FIELDS,
        },
        "columns": headers,
    }
    if preview:
        return summary
    if not valid:
        raise ValueError("No valid rows; existing dataset preserved")
    db.execute(update(Import).values(active=False))
    imp = Import(
        digest=digest,
        filename=filename.rsplit("/", 1)[-1][:255],
        actor=actor,
        summary=summary,
    )
    db.add(imp)
    db.flush()
    db.execute(update(Attendee).values(active=False))
    for gid, em, status, raw, norm in valid:
        db.add(ImportRecord(import_id=imp.id, guest_id=gid, raw=raw, normalized=norm))
        a = db.scalar(select(Attendee).where(Attendee.guest_id == gid))
        if not a:
            a = Attendee(guest_id=gid)
            db.add(a)
        a.email = em
        a.approval = status
        a.raw = raw
        a.normalized = norm
        a.active = True
        a.import_id = imp.id
    db.flush()
    return {"id": imp.id, "duplicate_import": False, **summary}
