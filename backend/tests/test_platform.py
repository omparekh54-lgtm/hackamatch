import io, csv, json, copy, zipfile, re, statistics, time
from pathlib import Path
from datetime import timedelta
from concurrent.futures import ThreadPoolExecutor
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select, func
from app.db import Session, now, engine
from app.models import *
from app.main import app
from app.auth import hash_password, issue, digest
from app.importer import import_data, normalize, email, FIELDS, parse
from app.service import eligibility, run_matching, publish, pool_for
from app.matching import allocate, score_team, WEIGHTS
from app.worker import tick

HEADERS = json.loads((Path(__file__).parent / "headers.json").read_text())


def rows(n=9):
    result = []
    for i in range(n):
        row = {h: "" for h in HEADERS}
        row.update(
            {
                "guest_id": f"synthetic-{i}",
                "email": f"person{i}@example.test",
                "name": f"Synthetic Person {i}",
                "approval_status": "approved",
                FIELDS["skills"]: [
                    "Backend",
                    "Front End, UI/UX",
                    "Presenting / public speaking",
                ][i % 3],
                FIELDS["goals"]: "Learn something new, Build a portfolio project",
                FIELDS["stretch"]: ["Strength", "Stretch", "Mix"][i % 3],
                FIELDS["experience"]: ["Never", "1–2", "3+"][i % 3],
                FIELDS["schedule"]: "Morning, Late night",
                FIELDS["conflict"]: [
                    "Debate it out",
                    "Vote",
                    "Defer to whoever owns that part",
                ][i % 3],
                FIELDS["leadership"]: ["Usually", "Sometimes", "Rarely"][i % 3],
                FIELDS["availability"]: "Yes",
                FIELDS["year"]: "3rd",
                FIELDS["hobbies"]: "reading, music",
                FIELDS["description"]: "Curious",
                FIELDS["feed"]: "Tech",
                FIELDS["soundtrack"]: "Lo-fi beats",
                FIELDS["college"]: "Synthetic College",
            }
        )
        result.append(row)
    return result


def csv_bytes(data):
    f = io.StringIO()
    w = csv.DictWriter(f, HEADERS)
    w.writeheader()
    w.writerows(data)
    return f.getvalue().encode()


def seed(n=9, verified=True):
    data = rows(n)
    with Session.begin() as db:
        import_data(db, csv_bytes(data), "synthetic.csv", None)
        w = Window(
            name="Synthetic short window",
            opens=now() - timedelta(minutes=1),
            closes=now() + timedelta(seconds=15),
            status="open",
            policy={
                "required_fields": ["skills", "goals", "schedule"],
                "require_offline": False,
            },
            config={"seed": 42, "weights": WEIGHTS},
        )
        db.add(w)
        db.flush()
        db.add(Job(window_id=w.id, key=f"match:{w.id}"))
        hashed = hash_password("Synthetic-pass-938!")
        ids = []
        for i in range(n):
            u = User(email=data[i]["email"], password=hashed, verified=verified)
            db.add(u)
            db.flush()
            db.add(Profile(user_id=u.id))
            db.add(Registration(user_id=u.id, window_id=w.id))
            ids.append(u.id)
        wid = w.id
    return wid, ids


def expire(wid):
    with Session.begin() as db:
        db.get(Window, wid).closes = now() - timedelta(seconds=1)


def as_user(uid):
    with Session.begin() as db:
        raw = issue(db, db.get(User, uid), "session", 24)
    c = TestClient(app)
    c.cookies.set("hm_session", raw)
    return c


def test_exact_header_import_and_idempotency():
    assert len(HEADERS) == 52
    with Session.begin() as db:
        r = import_data(db, csv_bytes(rows()), "synthetic.csv", None)
        assert r["accepted"] == 9
        assert len(r["columns"]) == 52
        again = import_data(db, csv_bytes(rows()), "same.csv", None)
        assert again["duplicate_import"]
        assert db.scalar(select(func.count()).select_from(Attendee)) == 9
        assert db.scalar(select(func.count()).select_from(Import)) == 1


def test_import_reconciliation_and_duplicate_email():
    data = rows()
    data[0]["email"] = "  PERSON0@EXAMPLE.TEST "
    data[1]["email"] = data[0]["email"]
    data[2]["approval_status"] = "declined"
    data[3][FIELDS["availability"]] = ""
    data.append(data[4].copy())
    bad = data[5].copy()
    bad["email"] = "not-an-email"
    bad["guest_id"] = "synthetic-invalid"
    data.append(bad)
    with Session.begin() as db:
        r = import_data(db, csv_bytes(data), "synthetic.csv", None)
        assert (
            r["accepted"] == 9
            and len(r["duplicates"]) == 1
            and len(r["rejected"]) == 1
            and r["ambiguous_emails"] == 1
        )
        u = User(email="person0@example.test", password="unused", verified=True)
        db.add(u)
        db.flush()
        assert eligibility(db, u, {})["state"] == "manual_review"
        import_data(db, csv_bytes(rows(2)), "new.csv", None)
        assert (
            db.scalar(
                select(func.count())
                .select_from(Attendee)
                .where(Attendee.active == True)
            )
            == 2
        )
        assert (
            db.scalar(
                select(func.count()).select_from(Import).where(Import.active == True)
            )
            == 1
        )


@pytest.mark.parametrize("value", ["", "not-email", "a@b", "a b@c.com"])
def test_invalid_emails(value):
    with pytest.raises(ValueError):
        email(value)


def test_normalization_unknown_not_negative():
    raw = rows(1)[0]
    raw[FIELDS["availability"]] = ""
    raw[FIELDS["skills"]] = ""
    raw[FIELDS["experience"]] = ""
    raw[FIELDS["leadership"]] = ""
    p = normalize(raw)
    assert (
        p["availability"] == "unknown"
        and p["experience"] is None
        and p["leadership"] is None
        and p["skills"] == []
    )
    assert normalize({**raw, FIELDS["year"]: "15"})["year"] is None
    assert normalize({**raw, FIELDS["skills"]: "Full stack"})["skills"] == ["fullstack"]
    assert email(" A@EXAMPLE.TEST ") == "a@example.test"


def test_unknown_missing_policy_and_declined():
    wid, ids = seed()
    with Session.begin() as db:
        u = db.get(User, ids[0])
        w = db.get(Window, wid)
        a = db.scalar(select(Attendee).where(Attendee.email == u.email))
        a.normalized = {**a.normalized, "availability": "unknown"}
        db.flush()
        assert eligibility(db, u, w.policy)["eligible"]
        assert not eligibility(db, u, {**w.policy, "require_offline": True})["eligible"]
        a.approval = "declined"
        db.flush()
        assert eligibility(db, u, w.policy)["state"] == "declined"
        a.approval = "approved"
        a.normalized = {**a.normalized, "skills": []}
        db.flush()
        assert not eligibility(db, u, w.policy)["eligible"]
        u.verified = False
        assert eligibility(db, u, w.policy)["state"] == "pending_verification"


def test_auth_flow_ownership_and_single_use(client):
    payload = {"email": " NEW@EXAMPLE.TEST ", "password": "Synthetic-pass-938!"}
    assert client.post("/api/auth/signup", json=payload).status_code == 200
    with Session() as db:
        u = db.scalar(select(User).where(User.email == "new@example.test"))
        assert not u.verified
        mail = db.scalar(select(Mail))
        token = mail.body.split("verify=")[1]
        assert db.get(Token, digest(token))
    assert client.post("/api/auth/login", json=payload).status_code == 200
    assert (
        client.get("/api/participant/profile").json()["state"] == "pending_verification"
    )
    assert client.post("/api/auth/verify", json={"token": token}).status_code == 200
    assert client.post("/api/auth/verify", json={"token": token}).status_code == 400
    assert client.get("/api/auth/me").json()["verified"]
    assert client.post("/api/auth/logout").status_code == 200
    assert client.get("/api/auth/me").status_code == 401
    assert (
        client.post(
            "/api/auth/login", json={**payload, "password": "incorrect-password"}
        ).status_code
        == 401
    )


def test_organizer_only_and_origin_protection(client):
    assert client.get("/api/admin/participants").status_code == 401
    wid, ids = seed()
    p = as_user(ids[0])
    assert p.get("/api/admin/participants").status_code == 403
    assert (
        p.post(
            "/api/participant/join", headers={"Origin": "https://evil.example.test"}
        ).status_code
        == 403
    )
    assert p.get("/api/admin/export").status_code == 403


def test_expired_verification_rejected(client):
    client.post(
        "/api/auth/signup",
        json={"email": "new@example.test", "password": "Synthetic-pass-938!"},
    )
    with Session.begin() as db:
        mail = db.scalar(select(Mail))
        raw = mail.body.split("verify=")[1]
        db.get(Token, digest(raw)).expires = now() - timedelta(seconds=1)
    assert client.post("/api/auth/verify", json={"token": raw}).status_code == 400


@pytest.mark.parametrize(
    "n,teams,leftovers",
    [(0, 0, 0), (1, 0, 1), (2, 0, 2), (9, 3, 0), (10, 3, 1), (20, 6, 2), (96, 32, 0)],
)
def test_matching_pool_sizes_seed_and_constraints(n, teams, leftovers):
    p = [{"id": i, **normalize(r)} for i, r in enumerate(rows(n))]
    r = allocate(p)
    s = allocate(list(reversed(p)))
    assert len(r["teams"]) == teams and len(r["unassigned"]) == leftovers
    assert r["teams"] == s["teams"] and r["unassigned"] == s["unassigned"]
    ids = [uid for t in r["teams"] for uid in t["members"]]
    assert len(ids) == len(set(ids)) and set(ids + r["unassigned"]) == set(range(n))
    assert all(len(t["members"]) == 3 for t in r["teams"])
    assert r["objective"] >= r["baseline"]
    assert r["runtime_ms"] < 10000


def test_missing_signals_and_unbalanced_skills():
    pool = [{"id": i, "skills": ["backend"], "experience": None} for i in range(20)]
    r = allocate(pool)
    assert len(r["teams"]) == 6 and len(r["unassigned"]) == 2
    breakdown = score_team([{}, {}, {}])
    assert breakdown["total"] == 50 and breakdown["components"]["skills"] is None
    with pytest.raises(ValueError):
        allocate([{"id": 1}, {"id": 1}])


def test_sensitive_fields_do_not_influence_matching():
    a = [{"id": i, **normalize(r)} for i, r in enumerate(rows(9))]
    b = [
        {
            **p,
            "gender": "other",
            "email": "any@example.test",
            "phone": "private",
            "sap": "private",
        }
        for p in a
    ]
    assert allocate(a)["teams"] == allocate(b)["teams"]


def test_background_restart_idempotency_and_publication(organizer):
    wid, ids = seed(10)
    p = as_user(ids[0])
    assert p.get("/api/participant/team").json()["team"] is None
    assert organizer.post(f"/api/admin/matching-windows/{wid}/start").status_code == 400
    expire(wid)
    # A fresh scheduler iteration reads persisted deadline after a simulated restart.
    tick()
    tick()
    r = organizer.post(f"/api/admin/matching-windows/{wid}/start")
    assert r.status_code == 200
    with Session() as db:
        assert db.get(Window, wid).status == "review"
        assert db.scalar(select(func.count()).select_from(Team)) == 3
        assert db.scalar(select(func.count()).select_from(Member)) == 9
        assert db.scalar(select(func.count()).select_from(Job)) == 1
        assert db.scalar(select(func.count()).select_from(Snapshot)) == 1
        assert (
            db.scalar(
                select(func.count())
                .select_from(Registration)
                .where(Registration.state == "waitlisted")
            )
            == 1
        )
    assert p.get("/api/participant/team").json()["team"] is None
    assert (
        organizer.post(f"/api/admin/matching-windows/{wid}/publish").status_code == 200
    )
    statuses = [as_user(uid).get("/api/participant/team").json() for uid in ids]
    assert sum(s["status"] == "assigned" for s in statuses) == 9
    assert sum(s["status"] == "waitlisted" for s in statuses) == 1
    team = next(s["team"] for s in statuses if s["team"])
    serialized = json.dumps(team)
    assert (
        "email" not in serialized
        and "approval_status" not in serialized
        and len(team["members"]) == 3
    )
    assert any(
        a["action"] == "teams.published"
        for a in organizer.get("/api/admin/audit").json()
    )
    r = organizer.get("/api/admin/export?format=csv")
    assert (
        r.status_code == 200
        and len(list(csv.reader(io.StringIO(r.content.decode("utf-8-sig"))))) == 11
    )
    r = organizer.get("/api/admin/export?format=xlsx")
    assert r.status_code == 200
    from openpyxl import load_workbook

    assert load_workbook(io.BytesIO(r.content)).active.max_row == 11


def test_publication_rechecks_current_approval(organizer):
    wid, ids = seed()
    expire(wid)
    tick()
    with Session.begin() as db:
        db.scalar(
            select(Attendee).where(Attendee.email == "person0@example.test")
        ).approval = "declined"
    assert (
        organizer.post(f"/api/admin/matching-windows/{wid}/publish").status_code == 400
    )


def test_join_window_deadline_and_unverified(client):
    wid, ids = seed()
    u = as_user(ids[0])
    assert u.post("/api/participant/join").status_code == 200
    with Session.begin() as db:
        db.get(User, ids[0]).verified = False
    assert u.post("/api/participant/join").status_code == 403
    expire(wid)
    assert as_user(ids[1]).post("/api/participant/join").status_code == 409


def test_window_configuration_and_one_active(organizer):
    r = organizer.post(
        "/api/admin/matching-windows",
        json={"name": "Synthetic round", "duration_seconds": 10},
    )
    assert r.status_code == 200
    wid = r.json()["id"]
    assert (
        organizer.post(
            "/api/admin/matching-windows", json={"name": "Another round"}
        ).status_code
        == 400
    )
    r = organizer.patch(
        f"/api/admin/matching-windows/{wid}",
        json={
            "reason": "Require offline attendance",
            "policy": {"require_offline": True},
            "config": {"weights": WEIGHTS, "seed": 11},
        },
    )
    assert r.status_code == 200 and r.json()["policy"]["require_offline"]
    assert (
        organizer.patch(
            f"/api/admin/matching-windows/{wid}",
            json={
                "reason": "Invalid weights test",
                "policy": {},
                "config": {"weights": {"skills": 101}},
            },
        ).status_code
        == 422
    )


def test_import_preview_authorization_and_xlsx(organizer):
    data = csv_bytes(rows())
    r = organizer.post(
        "/api/admin/luma/import?preview=true",
        files={"file": ("synthetic.csv", data, "text/csv")},
    )
    assert r.status_code == 200 and r.json()["accepted"] == 9
    assert organizer.get("/api/admin/luma/imports").json() == []
    r = organizer.post(
        "/api/admin/luma/import", files={"file": ("synthetic.csv", data, "text/csv")}
    )
    assert r.status_code == 200
    assert organizer.post(
        "/api/admin/luma/import", files={"file": ("synthetic.csv", data, "text/csv")}
    ).json()["duplicate_import"]
    assert (
        organizer.post(
            "/api/admin/luma/import",
            files={"file": ("bad.csv", b"wrong,headers\n1,2", "text/csv")},
        ).status_code
        == 400
    )
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(HEADERS)
    for r in rows(3):
        ws.append([r[h] for h in HEADERS])
    buf = io.BytesIO()
    wb.save(buf)
    assert parse(buf.getvalue(), "synthetic.xlsx")[0] == HEADERS


def test_overrides_survive_reimport(organizer):
    wid, ids = seed()
    uid = ids[0]
    assert (
        organizer.patch(
            f"/api/admin/participants/{uid}",
            json={"reason": "Confirmed with participant", "availability": "no"},
        ).status_code
        == 200
    )
    changed = rows()
    changed[0][FIELDS["availability"]] = "Yes"
    changed[0][FIELDS["hobbies"]] = "New hobby"
    with Session.begin() as db:
        import_data(db, csv_bytes(changed), "update.csv", None)
        e = eligibility(db, db.get(User, uid), {"require_offline": True})
        assert e["profile"]["availability"] == "no" and not e["eligible"]


def test_reassign_swap_waitlist_regenerate_and_reopen(organizer):
    wid, ids = seed(10)
    expire(wid)
    tick()
    teams = organizer.get("/api/admin/teams").json()
    out = teams[0]["members"][0]["id"]
    wait = organizer.get("/api/admin/unassigned").json()[0]["user_id"]
    r = organizer.post(
        f"/api/admin/teams/{teams[0]['id']}/reassign",
        json={
            "outgoing_user_id": out,
            "incoming_user_id": wait,
            "reason": "Balance team learning support",
        },
    )
    assert r.status_code == 200
    assert organizer.get("/api/admin/unassigned").json()[0]["user_id"] == out
    other = teams[1]["members"][0]["id"]
    assert (
        organizer.post(
            f"/api/admin/teams/{teams[0]['id']}/reassign",
            json={
                "outgoing_user_id": wait,
                "incoming_user_id": other,
                "reason": "Organizer approved swap",
            },
        ).status_code
        == 200
    )
    assert (
        organizer.post(
            f"/api/admin/matching-windows/{wid}/regenerate",
            json={"reason": "Explore another valid allocation"},
        ).status_code
        == 200
    )
    assert len(organizer.get("/api/admin/teams").json()) == 3
    assert (
        organizer.post(
            f"/api/admin/matching-windows/{wid}/reopen",
            json={"reason": "Invite additional eligible participants"},
        ).status_code
        == 200
    )
    assert organizer.get("/api/admin/teams").json() == []
    with Session() as db:
        assert db.scalar(select(func.count()).select_from(Snapshot)) == 0


def test_upload_limits_and_rate_limit(organizer, client):
    assert (
        organizer.post(
            "/api/admin/luma/import", files={"file": ("bad.txt", b"abc", "text/plain")}
        ).status_code
        == 400
    )
    p = {"email": "new@example.test", "password": "Synthetic-pass-938!"}
    for _ in range(3):
        assert client.post("/api/auth/signup", json=p).status_code == 200
    assert client.post("/api/auth/signup", json=p).status_code == 429


def test_sql_uniqueness_constrains_membership():
    from sqlalchemy.exc import IntegrityError

    wid, ids = seed()
    expire(wid)
    tick()
    with pytest.raises(IntegrityError):
        with Session.begin() as db:
            member = db.scalar(select(Member))
            db.add(
                Member(user_id=member.user_id, window_id=wid, team_id=member.team_id)
            )
            db.flush()


@pytest.mark.skipif(
    engine.dialect.name != "postgresql",
    reason="Requires TEST PostgreSQL database for row locks",
)
def test_concurrent_matching_requests():
    wid, ids = seed(20)
    expire(wid)

    def run(_):
        with Session.begin() as db:
            run_matching(db, wid)

    with ThreadPoolExecutor(max_workers=4) as pool:
        list(pool.map(run, range(4)))
    with Session() as db:
        assert db.scalar(select(func.count()).select_from(Team)) == 6
        assert db.scalar(select(func.count()).select_from(Member)) == 18
        assert db.scalar(select(func.count()).select_from(Job)) == 1


def test_conflicting_duplicate_id_is_quarantined():
    data = rows(3)
    extra = data[0].copy()
    extra["approval_status"] = "declined"
    data.append(extra)
    with Session.begin() as db:
        result = import_data(db, csv_bytes(data), "synthetic-conflict.csv", None)
        assert result["accepted"] == 2 and len(result["rejected"]) == 2
        assert not db.scalar(select(Attendee).where(Attendee.guest_id == "synthetic-0"))


def test_import_history_preserves_original_responses():
    original = rows(3)
    changed = rows(3)
    changed[0][FIELDS["skills"]] = "Research"
    with Session.begin() as db:
        first = import_data(db, csv_bytes(original), "v1.csv", None)
        import_data(db, csv_bytes(changed), "v2.csv", None)
        record = db.scalar(
            select(ImportRecord).where(
                ImportRecord.import_id == first["id"],
                ImportRecord.guest_id == "synthetic-0",
            )
        )
        assert record.raw[FIELDS["skills"]] == "Backend"
        assert db.scalar(
            select(Attendee).where(Attendee.guest_id == "synthetic-0")
        ).normalized["skills"] == ["research"]


def test_real_smtp_and_elapsed_short_window(organizer, monkeypatch):
    """Loopback SMTP, actual ten-second deadline, independent worker process."""
    import socket, subprocess, sys, os
    from aiosmtpd.controller import Controller

    class Inbox:
        def __init__(self):
            self.messages = []

        async def handle_DATA(self, server, session, envelope):
            self.messages.append((envelope.rcpt_tos, envelope.content.decode()))
            return "250 OK"

    inbox = Inbox()
    with socket.socket() as s:
        s.bind(("127.0.0.1", 0))
        port = s.getsockname()[1]
    smtp = Controller(inbox, hostname="127.0.0.1", port=port)
    smtp.start()
    monkeypatch.setenv("SMTP_HOST", "127.0.0.1")
    monkeypatch.setenv("SMTP_PORT", str(port))
    monkeypatch.setenv("SMTP_TLS", "false")
    clients = []
    try:
        assert (
            organizer.post(
                "/api/admin/luma/import",
                files={"file": ("synthetic.csv", csv_bytes(rows(9)), "text/csv")},
            ).status_code
            == 200
        )
        for i in range(9):
            c = TestClient(app)
            payload = {
                "email": f"person{i}@example.test",
                "password": "Synthetic-pass-938!",
            }
            assert c.post("/api/auth/signup", json=payload).status_code == 200
            clients.append(c)
        tick()
        assert len(inbox.messages) == 9
        for i, c in enumerate(clients):
            message = next(
                body
                for recipients, body in inbox.messages
                if f"person{i}@example.test" in recipients
            )
            # Email MIME may fold the URL; decode it before extracting the token.
            from email import policy
            from email.parser import Parser

            parsed = Parser(policy=policy.default).parsestr(message)
            token = re.search(r"verify=([A-Za-z0-9_-]+)", parsed.get_content()).group(1)
            assert c.post("/api/auth/verify", json={"token": token}).status_code == 200
            assert (
                c.post(
                    "/api/auth/login",
                    json={
                        "email": f"person{i}@example.test",
                        "password": "Synthetic-pass-938!",
                    },
                ).status_code
                == 200
            )
        w = organizer.post(
            "/api/admin/matching-windows",
            json={"name": "Elapsed-time integration test", "duration_seconds": 10},
        ).json()
        for c in clients:
            assert c.post("/api/participant/join").status_code == 200
        process = subprocess.Popen(
            [sys.executable, "-m", "app.worker"],
            env=os.environ.copy(),
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
        )
        try:
            deadline = time.monotonic() + 16
            while time.monotonic() < deadline:
                with Session() as db:
                    state = db.get(Window, w["id"]).status
                if state == "review":
                    break
                time.sleep(0.25)
            assert state == "review"
            assert (
                organizer.post(
                    f"/api/admin/matching-windows/{w['id']}/publish"
                ).status_code
                == 200
            )
            for c in clients:
                team = c.get("/api/participant/team").json()["team"]
                assert team and len(team["members"]) == 3
        finally:
            process.terminate()
            process.wait(timeout=5)
    finally:
        smtp.stop()


def test_scheduler_failure_rolls_back_and_retry_is_clean(monkeypatch):
    wid, ids = seed(9)
    expire(wid)
    import app.service as service

    original = service.allocate

    def broken(*a, **kw):
        raise RuntimeError("Synthetic failure")

    monkeypatch.setattr(service, "allocate", broken)
    tick()
    with Session() as db:
        assert db.get(Window, wid).status == "failed"
        assert db.scalar(select(func.count()).select_from(Team)) == 0
        assert db.scalar(select(func.count()).select_from(Member)) == 0
        assert db.scalar(select(func.count()).select_from(Snapshot)) == 0
    monkeypatch.setattr(service, "allocate", original)
    with Session.begin() as db:
        run_matching(db, wid)
    with Session() as db:
        assert db.get(Window, wid).status == "review"
        assert db.scalar(select(func.count()).select_from(Team)) == 3


def test_export_formula_prefix_is_neutralized(organizer):
    wid, ids = seed(3)
    expire(wid)
    tick()
    with Session.begin() as db:
        a = db.scalar(select(Attendee).where(Attendee.email == "person0@example.test"))
        a.raw = {**a.raw, "name": "=SYNTHETIC_FORMULA()"}
    assert (
        organizer.post(f"/api/admin/matching-windows/{wid}/publish").status_code == 200
    )
    r = organizer.get("/api/admin/export?format=csv")
    assert "'=SYNTHETIC_FORMULA()" in r.content.decode("utf-8-sig")


def test_unrecognized_skills_do_not_become_zero_ability_scores():
    team = [{"skills": ["other"], "stretch": "stretch"} for _ in range(3)]
    score = score_team(team)
    assert score["components"]["skills"] is None
    assert score["components"]["roles"] is None
    assert score["components"]["style"] is None
    assert score["total"] == 50
